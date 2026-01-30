import streamlit as st
import google.generativeai as genai
import markdown
import pandas as pd
import io
import docx
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from pyairtable import Api

# 1. SETUP & PAGE CONFIG
st.set_page_config(
    page_title="Story Grid Analyzer Pro",
    page_icon="🧬",
    layout="wide",
    initial_sidebar_state="expanded"
)

# --- 📚 PROJECT CONFIGURATION (THE SIMPLIFIED "ONE ID" MODEL) ---
PROJECTS = {
    "Chione Trilogy": {
        "icon": "❄️",
        "theme_color": "#008080", # Teal
        "locale": "American English",
        "base_id": "apphWWXcojkv7nPni", 
        "codex_table_id": "tblNNN5029XqoHdYV" # <--- JUST ONE MASTER TABLE ID!
    },
    "The Gatekeepers": {
        "icon": "🚪",
        "theme_color": "#D2691E", # Desert Orange
        "locale": "American English",
        "base_id": "app...",        # <--- Paste Real Base ID
        "codex_table_id": "tbl..."  # <--- Paste Real Master Table ID
    }
}

# --- CUSTOM CSS (DYNAMIC THEMING) ---
def apply_theme(color):
    st.markdown(f"""
    <style>
        #MainMenu {{visibility: hidden;}}
        footer {{visibility: hidden;}}
        header {{visibility: hidden;}}
        h1, h2, h3, h4 {{ color: {color} !important; }}
        div.stButton > button:first-child {{
            border: 2px solid {color};
            color: {color};
        }}
        div.stButton > button:first-child:hover {{
            background-color: {color};
            color: white;
        }}
        div[data-testid="stMetricValue"] {{ color: {color}; }}
        .stApp {{ background-image: linear-gradient(to bottom right, #ffffff, #f0f2f6); }}
    </style>
    """, unsafe_allow_html=True)

# --- SECURITY ---
def check_password():
    if "password_correct" not in st.session_state:
        st.session_state.password_correct = False
    if st.session_state.password_correct:
        return True
    st.title("🔒 Login Required")
    password = st.text_input("Enter Password", type="password")
    if st.button("Enter"):
        if password == "story2026": 
            st.session_state.password_correct = True
            st.rerun()
        else:
            st.error("😕 Incorrect password")
    return False

if not check_password():
    st.stop()

# --- SESSION STATE ---
if 'chapter_log' not in st.session_state: st.session_state.chapter_log = []
if 'current_report' not in st.session_state: st.session_state.current_report = ""

# --- API SETUP ---
try:
    genai.configure(api_key=st.secrets["GOOGLE_API_KEY"])
except:
    st.error("⚠️ API Key missing. Please check Secrets.")

# --- AIRTABLE FUNCTIONS (READ & WRITE) ---
@st.cache_data(ttl=10) # Fast refresh so we see new items immediately
def fetch_master_codex(base_id, table_id):
    try:
        if "AIRTABLE_TOKEN" in st.secrets:
            api = Api(st.secrets["AIRTABLE_TOKEN"])
            table = api.table(base_id, table_id)
            records = table.all()
            data = [r['fields'] for r in records]
            return pd.DataFrame(data)
        else:
            return None
    except Exception as e:
        if "REPLACE" in table_id: return pd.DataFrame({"Info": ["Config Needed"]})
        return None

def add_to_codex(base_id, table_id, name, category, description):
    try:
        api = Api(st.secrets["AIRTABLE_TOKEN"])
        table = api.table(base_id, table_id)
        table.create({"Name": name, "Category": category, "Description": description})
        return True
    except Exception as e:
        st.error(f"Failed to save: {e}")
        return False

# --- HELPERS ---
def read_file(uploaded_file):
    if uploaded_file is None: return ""
    if uploaded_file.name.endswith(".docx"):
        try:
            doc = docx.Document(uploaded_file)
            return "\n".join([para.text for para in doc.paragraphs])
        except: return ""
    elif uploaded_file.name.endswith(".txt"):
        try: return uploaded_file.read().decode("utf-8")
        except: return ""
    return ""

def create_html_report(content, title):
    html_content = markdown.markdown(content, extensions=['tables'])
    return f"<html><body><h1>{title}</h1>{html_content}</body></html>"

def clean_markdown_text(text):
    if not isinstance(text, str): return text
    return text.replace('**', '').replace('__', '').replace('### ', '').replace('## ', '').replace('# ', '')

def to_excel(df):
    export_df = df.copy()
    for col in export_df.columns:
        export_df[col] = export_df[col].apply(clean_markdown_text)
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        export_df.to_excel(writer, index=False, sheet_name='Story Map')
        worksheet = writer.sheets['Story Map']
        header_fill = PatternFill(start_color="008080", end_color="008080", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True, size=12, name="Calibri")
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
        for i, col in enumerate(export_df.columns):
            worksheet.column_dimensions[chr(65 + i)].width = 25
    return output.getvalue()

def analyze_outline(text, genre, framework, locale):
    prompt = f"Analyze this PLOT OUTLINE against {framework}. Output Structural Health Score, Gap Analysis, and Pacing Check. Use {locale} spelling.\n\nTEXT: {text}"
    model = genai.GenerativeModel('gemini-flash-latest')
    return model.generate_content(prompt).text

def analyze_scene(text, genre, framework, beat, locale):
    prompt = f"""
    You are a ruthless Story Grid editor. Analyze this SCENE.
    LANGUAGE SETTING: Write in strictly {locale}.
    STYLE: RUTHLESSLY CONCISE. Use Markdown Tables.
    CONTEXT: Genre: {genre} | Framework: {framework} | Beat: {beat}
    OUTPUT FORMAT:
    1. HEADER: "# [Emoji] [Start Value] ➔ [End Value]"
    2. THE 5 COMMANDMENTS (Bullet points)
    3. VISUAL SCOREBOARD (Markdown Table: Element, Start Value, End Value, Notes)
    """
    if "None" not in framework: prompt += f"\n4. FRAMEWORK CHECK ({framework}): Verdict & Reason."
    prompt += f"\n5. GENRE CHECK: Gold Star Moment 🌟\n\nSTORY TEXT: {text}"
    model = genai.GenerativeModel('gemini-flash-latest')
    return model.generate_content(prompt).text

# --- UI START ---

# 1. SIDEBAR NAVIGATION
with st.sidebar:
    st.title("Navigator")
    st.markdown("### 🗂️ Active Project")
    selected_project_key = st.selectbox("Select Book/Series", list(PROJECTS.keys()), label_visibility="collapsed")
    
    # Load Config
    config = PROJECTS[selected_project_key]
    apply_theme(config["theme_color"]) # Apply Theme
    
    st.divider()
    st.markdown("### ✍️ Creator Tools")
    
    # --- QUICK ADD FORM (The new "Write" Feature) ---
    with st.expander("➕ Quick Add to Codex", expanded=False):
        with st.form("add_record_form", clear_on_submit=True):
            new_name = st.text_input("Name", placeholder="e.g. The Crystal Key")
            new_cat = st.selectbox("Category", ["Character", "Location", "Lore", "Item", "Faction"])
            new_desc = st.text_area("Description", placeholder="A brief summary...")
            if st.form_submit_button("Save to Airtable"):
                if add_to_codex(config["base_id"], config["codex_table_id"], new_name, new_cat, new_desc):
                    st.toast(f"Saved {new_name}!", icon="💾")
                    fetch_master_codex.clear() # Force refresh data
                    st.rerun()

    st.divider()
    st.markdown("### ⚙️ Analysis Config")
    selected_genre = st.selectbox("Genre", ["Action/Thriller", "Love Story", "Horror", "Mystery/Crime", "Sci-Fi/Fantasy"])
    selected_framework = st.selectbox("Framework", ["None (Pure Story Grid)", "Save the Cat!", "Dan Harmon's Story Circle"])
    
    scene_count = len(st.session_state.chapter_log)
    st.metric("Scenes Logged", scene_count)
    if st.button("🗑️ Clear Session Data"):
        st.session_state.chapter_log = []
        st.rerun()

# 2. MAIN HEADER
col_h1, col_h2 = st.columns([1, 6])
with col_h1: st.title(config["icon"])
with col_h2:
    st.title(f"{selected_project_key}")
    st.caption(f"Story Grid Analyzer Pro • {selected_genre} • {config['locale']}")

m1, m2, m3, m4 = st.columns(4)
with m1: st.metric("World Status", "Online", "Connected")
with m2: st.metric("Codex Mode", "Unified", "Read/Write")
with m3: st.metric("Scenes Tracked", scene_count)
with m4: st.metric("Draft Health", "Pending")
st.markdown("---")

# 3. TABS
tab1, tab2, tab3 = st.tabs(["🩺 Outline Doctor", "🔬 Scene Logger", "📚 World Codex"])

# === TAB 1: OUTLINE DOCTOR ===
with tab1:
    c1, c2 = st.columns([2, 1])
    with c1:
        uploaded_outline = st.file_uploader("Drop your full outline or synopsis here", type=["docx", "txt"], key="outline")
        outline_text = read_file(uploaded_outline) if uploaded_outline else ""
    with c2: st.info("💡 Paste your beat sheet here for a check-up.")
    outline_input = st.text_area("Or paste text:", value=outline_text, height=300)
    
    if st.button("✨ Run Diagnosis", type="primary", use_container_width=True):
        if outline_input:
            with st.spinner(f"Analyzing..."):
                try:
                    res = analyze_outline(outline_input, selected_genre, selected_framework, config['locale'])
                    st.markdown(res)
                    st.balloons()
                except Exception as e: st.error(f"Error: {e}")

# === TAB 2: SCENE LOGGER ===
with tab2:
    col1, col2 = st.columns([2, 1])
    with col1:
        with st.container(border=True):
            chapter_title = st.text_input("Chapter Title", placeholder="e.g. Chapter 1")
            uploaded_scene = st.file_uploader("Upload Scene", type=["docx", "txt"], key="scene")
            scene_text = read_file(uploaded_scene) if uploaded_scene else ""
            scene_input = st.text_area("Scene Text", value=scene_text, height=250)
            
            if st.button("🚀 Analyze Scene", type="primary", use_container_width=True):
                if scene_input:
                    with st.spinner("Analyzing..."):
                        try:
                            rep = analyze_scene(scene_input, selected_genre, selected_framework, "General", config['locale'])
                            st.session_state.current_report = rep
                            st.toast("Done!", icon="✅")
                            st.rerun()
                        except Exception as e: st.error(f"Error: {e}")

    with col2:
        if st.session_state.current_report:
            if st.button("➕ Add to Log", use_container_width=True):
                st.session_state.chapter_log.append({"Chapter": chapter_title, "Analysis": st.session_state.current_report})
                st.toast("Saved!", icon="💾")
            st.download_button("📥 Download Report", create_html_report(st.session_state.current_report, chapter_title), "report.html", use_container_width=True)

    if st.session_state.current_report:
        with st.expander("📝 Full Report", expanded=True): st.markdown(st.session_state.current_report)

# === TAB 3: DYNAMIC WORLD CODEX ===
with tab3:
    # Fetch the ONE master table
    df_codex = fetch_master_codex(config["base_id"], config["codex_table_id"])
    
    if df_codex is not None and not df_codex.empty and "Category" in df_codex.columns:
        # Get unique categories (e.g., Character, Location)
        categories = df_codex["Category"].unique()
        
        # Create dynamic sub-tabs for each category
        sub_tabs = st.tabs(list(categories))
        
        for i, cat in enumerate(categories):
            with sub_tabs[i]:
                # Filter data for this tab
                df_filtered = df_codex[df_codex["Category"] == cat]
                
                # Search Bar specific to this category
                search = st.text_input(f"Search {cat}s...", key=f"search_{cat}")
                if search:
                    mask = df_filtered.apply(lambda x: x.astype(str).str.contains(search, case=False).any(), axis=1)
                    df_filtered = df_filtered[mask]
                
                st.dataframe(df_filtered, use_container_width=True, hide_index=True)
                st.caption(f"{len(df_filtered)} items")
    else:
        st.info("👋 Welcome! To start, ensure your Airtable has a column named 'Category'. Use the Sidebar to add your first item!")

# === FOOTER ===
st.divider()
if len(st.session_state.chapter_log) > 0:
    st.subheader("📊 Session Log")
    df = pd.DataFrame(st.session_state.chapter_log)
    st.dataframe(df, use_container_width=True)
    st.download_button("📥 Download Excel", to_excel(df), "log.xlsx")